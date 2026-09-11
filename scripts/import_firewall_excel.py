from __future__ import annotations

import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook


DEFAULT_CREATOR = "andyliu%40as.edu.tw"


@dataclass
class FirewallRow:
    legacy_form_number: str
    system_name: str
    action: str
    purpose_type: str
    source_zone: str
    source_zone2: str
    source_ip: str
    destination_zone: str
    destination_zone2: str
    destination_ip: str
    protocol_type: str
    start_date: str
    end_date: str
    request_date: str
    rule_description: str
    firewall_zone: str
    firewall_id: str


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r", "\n")
    text = re.sub(r"\s+", " ", text.strip())
    return text


def normalize_date(value: object) -> str:
    text = normalize_text(value)
    if not text:
        return ""
    text = text.replace("/", "-")
    parts = [part for part in text.split("-") if part]
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        year, month, day = parts
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return text


def dedupe_key(row: FirewallRow) -> tuple[str, ...]:
    return (
        normalize_text(row.legacy_form_number),
        normalize_text(row.action),
        normalize_text(row.source_zone),
        normalize_text(row.source_zone2),
        normalize_text(row.source_ip),
        normalize_text(row.destination_zone),
        normalize_text(row.destination_zone2),
        normalize_text(row.destination_ip),
        normalize_text(row.protocol_type),
        normalize_date(row.start_date),
        normalize_date(row.end_date),
        normalize_date(row.request_date),
        normalize_text(row.firewall_id),
    )


def load_rows(xlsx_path: Path) -> list[FirewallRow]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[FirewallRow] = []
    last_system_name = ""
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        values = list(row[:18])
        if not any(value not in (None, "") for value in values):
            continue
        system_name = normalize_text(values[2]) or last_system_name
        if system_name:
            last_system_name = system_name
        rows.append(
            FirewallRow(
                legacy_form_number=normalize_text(values[1]),
                system_name=system_name,
                action=normalize_text(values[3]),
                purpose_type=normalize_text(values[4]),
                source_zone=normalize_text(values[5]),
                source_zone2=normalize_text(values[6]),
                source_ip=normalize_text(values[7]),
                destination_zone=normalize_text(values[8]),
                destination_zone2=normalize_text(values[9]),
                destination_ip=normalize_text(values[10]),
                protocol_type=normalize_text(values[11]),
                start_date=normalize_date(values[12]),
                end_date=normalize_date(values[13]),
                request_date=normalize_date(values[14]),
                rule_description=normalize_text(values[15]),
                firewall_zone=normalize_text(values[16]),
                firewall_id=normalize_text(values[17]),
            )
        )
    return rows


def existing_key_ids(conn: sqlite3.Connection) -> dict[tuple[str, ...], list[int]]:
	cur = conn.cursor()
	result: dict[tuple[str, ...], list[int]] = {}
	for db_row in cur.execute(
		"""
        SELECT id, legacy_form_number, system_name, action, purpose_type, source_zone, source_zone2,
               source_ip, destination_zone, destination_zone2, destination_ip, protocol_type,
               start_date, end_date, request_date, rule_description, firewall_zone, firewall_id
        FROM firewall_requests
        """
	):
		row = FirewallRow(
			legacy_form_number=db_row[1],
			system_name=db_row[2],
			action=db_row[3],
			purpose_type=db_row[4],
			source_zone=db_row[5],
			source_zone2=db_row[6],
			source_ip=db_row[7],
			destination_zone=db_row[8],
			destination_zone2=db_row[9],
			destination_ip=db_row[10],
			protocol_type=db_row[11],
			start_date=db_row[12],
			end_date=db_row[13],
			request_date=db_row[14],
			rule_description=db_row[15],
			firewall_zone=db_row[16],
			firewall_id=db_row[17],
		)
		result.setdefault(dedupe_key(row), []).append(db_row[0])
	return result


def insert_rows(conn: sqlite3.Connection, rows: list[FirewallRow], creator: str) -> tuple[int, int, int]:
    unique_rows = {dedupe_key(row): row for row in rows}
    current_key_ids = existing_key_ids(conn)
    inserted = 0
    replaced = 0
    cur = conn.cursor()
    for key, ids in current_key_ids.items():
        if key not in unique_rows:
            continue
        for row_id in ids:
            cur.execute("DELETE FROM firewall_requests WHERE id=?", (row_id,))
            replaced += 1
    for row in unique_rows.values():
        cur.execute(
            """
            INSERT INTO firewall_requests (
                legacy_form_number, system_name, action, purpose_type, source_zone, source_zone2,
                source_ip, destination_zone, destination_zone2, destination_ip, protocol_type,
                start_date, end_date, request_date, rule_description, firewall_zone, firewall_id,
                status, creator, remarks
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, '')
            """,
            (
                row.legacy_form_number,
                row.system_name,
                row.action,
                row.purpose_type,
                row.source_zone,
                row.source_zone2,
                row.source_ip,
                row.destination_zone,
                row.destination_zone2,
                row.destination_ip,
                row.protocol_type,
                row.start_date,
                row.end_date,
                row.request_date,
                row.rule_description,
                row.firewall_zone,
                row.firewall_id,
                creator,
            ),
        )
        inserted += 1
    conn.commit()
    return inserted, replaced, len(rows) - len(unique_rows)


def main() -> int:
    if len(sys.argv) < 3:
        print("usage: python scripts/import_firewall_excel.py <xlsx_path> <db_path> [creator]")
        return 1
    xlsx_path = Path(sys.argv[1])
    db_path = Path(sys.argv[2])
    creator = sys.argv[3] if len(sys.argv) > 3 else DEFAULT_CREATOR
    rows = load_rows(xlsx_path)
    conn = sqlite3.connect(db_path)
    try:
        inserted, replaced, excel_duplicates = insert_rows(conn, rows, creator)
    finally:
        conn.close()
    print(f"rows={len(rows)} inserted={inserted} replaced={replaced} excel_duplicates={excel_duplicates} creator={creator}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
